from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Iterable

from .contracts import SourceLocator, build_projection_bundle, make_native_ref
from .docling_source import DoclingSourceError, DoclingSourceUnavailable, prepare_docling_source
from .source_contract import NEO_SOURCE_DOCUMENT_SCHEMA_ID, build_source_document, file_sha256

PROJECT_DOCUMENT_EXTRACTION_SCHEMA_ID = NEO_SOURCE_DOCUMENT_SCHEMA_ID  # compatibility alias
PROJECT_DOCUMENT_FRAGMENT_SCHEMA_ID = "neo.knowledge.project_document_fragment.v1"
PROJECT_DOCUMENT_MANIFEST_SCHEMA_ID = "neo.knowledge.project_document_manifest.v1"
PROJECT_DOCUMENT_ADAPTER_ID = "neo.project_documents"
ADVANCED_DOCUMENT_SUFFIXES = {".pdf", ".docx", ".pptx", ".xlsx"}

# Persistent project files are bounded by upload size, not an arbitrary text preview
# ceiling. Fragment sizes are deliberately small because these are retrieval units,
# not prompt payloads.
TARGET_FRAGMENT_CHARS = 2_800
MAX_FRAGMENT_CHARS = 4_200
LONG_BLOCK_OVERLAP_CHARS = 240

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
_KEY_VALUE_RE = re.compile(r"^\s*([A-Za-z][A-Za-z0-9 _/\-]{1,48})\s*:\s*(.+?)\s*$")
_EXPLICIT_ENTITY_RE = re.compile(
    r"^\s*(Character|Person|Location|Place|Organization|Company|Project|Client|Product|Model|Team|Scene|Episode)\s*:\s*(.+?)\s*$",
    re.I,
)

_ENTITY_SECTION_TYPES: tuple[tuple[str, str], ...] = (
    ("characters", "character"), ("character", "character"), ("people", "person"), ("persons", "person"),
    ("locations", "location"), ("places", "location"), ("organizations", "organization"),
    ("companies", "organization"), ("projects", "project"), ("clients", "client"),
    ("products", "product"), ("models", "model"), ("teams", "team"),
)
_ALIAS_KEYS = {"alias", "aliases", "aka", "also known as", "nickname", "nicknames"}
_RELATION_KEYS = {
    "brother", "sister", "sibling", "father", "mother", "parent", "child", "son", "daughter",
    "boss", "manager", "works for", "works_for", "friend", "partner", "spouse", "lover", "mentor",
    "student", "assistant", "secretary", "owner", "member of", "member_of",
}

_UNRESOLVED_HEADING_MARKERS = (
    "open question", "open questions", "questions preserved", "unresolved", "unknown",
    "questions remaining", "questions remain", "what remains unknown", "archive questions",
)
_SPECULATIVE_HEADING_MARKERS = (
    "possibilities", "possible reading", "possible readings", "interpretations", "theories",
    "provisional", "speculation", "speculative", "candidate reading", "compatible origin readings",
)


def classify_fragment_epistemics(*, title: str = "", heading_path: list[str] | None = None, text: str = "") -> dict[str, Any]:
    """Classify what a Project passage can establish, without reinterpreting its content.

    A source passage can itself be verified while the claims inside it are explicitly
    unresolved/speculative. Retrieval and grounding need that distinction so a
    question such as ``Who discovered the final seal?`` cannot become evidence that
    a nearby entity *did* discover it.
    """
    headings = [str(item or "").strip() for item in (heading_path or []) if str(item or "").strip()]
    heading_text = " > ".join(headings)
    title_text = str(title or "")
    body = str(text or "").strip()
    scope_text = f"{heading_text} {title_text}".lower()

    role = "declarative_passage"
    state = "established"
    claim_type = "source_passage"
    supports_positive_claims = True
    reason = "default_declarative_source_passage"

    if any(marker in scope_text for marker in _UNRESOLVED_HEADING_MARKERS):
        role = "unresolved_question"
        state = "explicit_unknown"
        claim_type = "unresolved_question"
        supports_positive_claims = False
        reason = "unresolved_heading"
    elif any(marker in scope_text for marker in _SPECULATIVE_HEADING_MARKERS):
        role = "speculative_passage"
        state = "inferred"
        claim_type = "speculative_source_passage"
        supports_positive_claims = False
        reason = "speculative_heading"
    else:
        nonempty = [line.strip() for line in body.splitlines() if line.strip()]
        questions = [line for line in nonempty if line.rstrip().endswith("?")]
        declarative = [line for line in nonempty if not line.rstrip().endswith("?") and not re.match(r"^[-*+]\s*$", line)]
        # Question-list sections without an explicit heading still represent
        # unresolved prompts, not affirmative evidence. Keep the threshold
        # conservative so ordinary prose containing one rhetorical question is
        # not downgraded.
        if len(questions) >= 3 and len(questions) >= max(3, len(declarative) * 2):
            role = "unresolved_question"
            state = "explicit_unknown"
            claim_type = "unresolved_question"
            supports_positive_claims = False
            reason = "question_list_shape"

    return {
        "epistemic_role": role,
        "epistemic_state": state,
        "claim_type": claim_type,
        "supports_positive_claims": supports_positive_claims,
        "classification_reason": reason,
    }


def _hash(value: Any, length: int = 32) -> str:
    if isinstance(value, str):
        text = value
    else:
        text = json.dumps(value, ensure_ascii=False, sort_keys=True, default=str, separators=(",", ":"))
    return hashlib.sha256(text.encode("utf-8", errors="ignore")).hexdigest()[:length]


def _clean(text: Any) -> str:
    value = str(text or "").replace("\x00", "")
    value = re.sub(r"\r\n?", "\n", value)
    value = re.sub(r"[ \t]+\n", "\n", value)
    return value.strip()


def _slug(text: str) -> str:
    clean = re.sub(r"[^a-z0-9]+", "_", str(text or "").lower()).strip("_")
    return clean[:80] or "value"


def _block(*, kind: str, text: str, heading_path: list[str] | None = None, locator: dict[str, Any] | None = None) -> dict[str, Any] | None:
    body = _clean(text)
    if not body:
        return None
    loc = dict(locator or {})
    block_id = f"blk_{_hash({'kind': kind, 'locator': loc, 'text': body}, 20)}"
    return {
        "block_id": block_id,
        "kind": str(kind or "paragraph"),
        "text": body,
        "heading_path": [str(item) for item in (heading_path or []) if str(item or "").strip()],
        "locator": loc,
        "char_count": len(body),
    }


def _heading_update(stack: list[str], level: int, title: str) -> list[str]:
    level = max(1, min(int(level or 1), 8))
    next_stack = list(stack[: level - 1])
    while len(next_stack) < level - 1:
        next_stack.append("")
    next_stack.append(_clean(title))
    return [item for item in next_stack if item]


def _blocks_from_lines(text: str, *, markdown: bool = False, source_kind: str = "text") -> list[dict[str, Any]]:
    lines = _clean(text).splitlines()
    blocks: list[dict[str, Any]] = []
    heading_path: list[str] = []
    buffer: list[str] = []
    start_line = 1

    def flush(end_line: int) -> None:
        nonlocal buffer, start_line
        body = "\n".join(buffer).strip()
        if body:
            item = _block(kind="paragraph", text=body, heading_path=heading_path, locator={"kind": "line_range", "start_line": start_line, "end_line": end_line, "source_kind": source_kind})
            if item:
                blocks.append(item)
        buffer = []

    for idx, line in enumerate(lines, start=1):
        match = _HEADING_RE.match(line.strip()) if markdown else None
        if match:
            flush(idx - 1)
            heading_path = _heading_update(heading_path, len(match.group(1)), match.group(2))
            blocks.append(_block(kind="heading", text=match.group(2), heading_path=heading_path, locator={"kind": "line", "line": idx, "source_kind": source_kind}) or {})
            start_line = idx + 1
            continue
        if not line.strip():
            flush(idx - 1)
            start_line = idx + 1
            continue
        if not buffer:
            start_line = idx
        buffer.append(line)
    flush(len(lines))
    return [item for item in blocks if item and item.get("text")]


class _StructuredHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[dict[str, Any]] = []
        self.heading_path: list[str] = []
        self.capture_tag = ""
        self.capture: list[str] = []
        self.index = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        if tag in {"h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "td", "th"}:
            self.capture_tag = tag
            self.capture = []

    def handle_data(self, data: str) -> None:
        if self.capture_tag:
            self.capture.append(data)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag != self.capture_tag:
            return
        text = _clean(" ".join(self.capture))
        self.index += 1
        if text:
            if tag.startswith("h") and len(tag) == 2 and tag[1].isdigit():
                self.heading_path = _heading_update(self.heading_path, int(tag[1]), text)
                item = _block(kind="heading", text=text, heading_path=self.heading_path, locator={"kind": "html_node", "node_index": self.index, "tag": tag})
            else:
                item = _block(kind="table_cell" if tag in {"td", "th"} else "paragraph", text=text, heading_path=self.heading_path, locator={"kind": "html_node", "node_index": self.index, "tag": tag})
            if item:
                self.blocks.append(item)
        self.capture_tag = ""
        self.capture = []


def _extract_text(path: Path, suffix: str) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    raw = path.read_bytes()
    decoded = ""
    encoding = "utf-8"
    for candidate in ("utf-8-sig", "utf-8", "utf-16", "latin-1"):
        try:
            decoded = raw.decode(candidate)
            encoding = candidate
            break
        except UnicodeDecodeError:
            continue
    if suffix in {".html", ".htm"}:
        parser = _StructuredHTMLParser(); parser.feed(decoded)
        return parser.blocks, {"method": "html_parser", "encoding": encoding}
    if suffix in {".json", ".jsonl"}:
        # Preserve source text exactly enough for citation while making top-level JSON
        # sections searchable when parsing succeeds.
        try:
            parsed = json.loads(decoded) if suffix == ".json" else None
        except Exception:
            parsed = None
        if isinstance(parsed, dict):
            blocks = []
            for idx, (key, value) in enumerate(parsed.items(), start=1):
                body = json.dumps(value, ensure_ascii=False, indent=2, default=str)
                item = _block(kind="json_section", text=f"{key}:\n{body}", heading_path=[str(key)], locator={"kind": "json_key", "key": str(key), "index": idx})
                if item: blocks.append(item)
            if blocks:
                return blocks, {"method": "json_top_level", "encoding": encoding}
    if suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        reader = csv.reader(io.StringIO(decoded), delimiter=delimiter)
        rows = list(reader)
        if rows:
            header = rows[0]
            blocks = []
            for start in range(1, len(rows), 40):
                chunk = rows[start:start + 40]
                sio = io.StringIO(); writer = csv.writer(sio, delimiter=delimiter, lineterminator="\n")
                writer.writerow(header)
                writer.writerows(chunk)
                item = _block(kind="table_rows", text=sio.getvalue(), heading_path=[path.stem], locator={"kind": "row_range", "start_row": start + 1, "end_row": start + len(chunk)})
                if item: blocks.append(item)
            return blocks, {"method": "delimited_rows", "encoding": encoding, "row_count": len(rows)}
    return _blocks_from_lines(decoded, markdown=suffix in {".md", ".markdown"}, source_kind=suffix.lstrip(".") or "text"), {"method": "structured_text", "encoding": encoding}


def _extract_pdf(path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    try:
        from pypdf import PdfReader  # type: ignore
    except Exception as exc:
        return [], {"status": "stored_only", "reason": f"pdf_text_extractor_unavailable: {exc}"}
    try:
        reader = PdfReader(str(path))
        blocks: list[dict[str, Any]] = []
        for page_number, page in enumerate(reader.pages, start=1):
            try:
                page_text = _clean(page.extract_text() or "")
            except Exception:
                page_text = ""
            if not page_text:
                continue
            paragraphs = [part.strip() for part in re.split(r"\n\s*\n", page_text) if part.strip()]
            if not paragraphs:
                paragraphs = [page_text]
            for paragraph_index, paragraph in enumerate(paragraphs, start=1):
                item = _block(kind="paragraph", text=paragraph, heading_path=[], locator={"kind": "pdf_page", "page": page_number, "paragraph": paragraph_index})
                if item: blocks.append(item)
        return blocks, {"method": "pypdf_all_pages", "page_count": len(reader.pages), "pages_scanned": len(reader.pages)}
    except Exception as exc:
        return [], {"status": "stored_only", "reason": f"pdf_extract_failed: {exc}"}


def _extract_docx(path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    try:
        from docx import Document  # type: ignore
        document = Document(str(path))
        blocks: list[dict[str, Any]] = []
        heading_path: list[str] = []
        paragraph_index = 0
        for paragraph in document.paragraphs:
            paragraph_index += 1
            text = _clean(paragraph.text)
            if not text:
                continue
            style_name = str(getattr(getattr(paragraph, "style", None), "name", "") or "")
            match = re.match(r"Heading\s+(\d+)", style_name, re.I)
            if match:
                heading_path = _heading_update(heading_path, int(match.group(1)), text)
                kind = "heading"
            else:
                kind = "paragraph"
            item = _block(kind=kind, text=text, heading_path=heading_path, locator={"kind": "docx_paragraph", "paragraph": paragraph_index, "style": style_name})
            if item: blocks.append(item)
        for table_index, table in enumerate(document.tables, start=1):
            for row_index, row in enumerate(table.rows, start=1):
                values = [_clean(cell.text) for cell in row.cells]
                text = " | ".join(value for value in values if value)
                item = _block(kind="table_row", text=text, heading_path=heading_path, locator={"kind": "docx_table_row", "table": table_index, "row": row_index})
                if item: blocks.append(item)
        return blocks, {"method": "python_docx", "paragraph_count": len(document.paragraphs), "table_count": len(document.tables)}
    except Exception:
        # Dependency-free fallback preserves paragraphs but may not recover styles.
        try:
            import zipfile
            import xml.etree.ElementTree as ET
            with zipfile.ZipFile(path) as zf:
                xml = zf.read("word/document.xml")
            root = ET.fromstring(xml)
            ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            blocks = []
            for idx, paragraph in enumerate(root.iter(f"{ns}p"), start=1):
                text = _clean("".join(node.text or "" for node in paragraph.iter(f"{ns}t")))
                item = _block(kind="paragraph", text=text, heading_path=[], locator={"kind": "docx_paragraph", "paragraph": idx})
                if item: blocks.append(item)
            return blocks, {"method": "docx_xml_fallback", "paragraph_count": len(blocks)}
        except Exception as exc:
            return [], {"status": "stored_only", "reason": f"docx_extract_failed: {exc}"}


def _extract_pptx(path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    try:
        from pptx import Presentation  # type: ignore
        presentation = Presentation(str(path))
        blocks = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            title = ""
            try:
                title = _clean(slide.shapes.title.text if slide.shapes.title else "")
            except Exception:
                title = ""
            heading = [title] if title else [f"Slide {slide_number}"]
            for shape_index, shape in enumerate(slide.shapes, start=1):
                text = _clean(getattr(shape, "text", "") or "")
                if not text or text == title:
                    continue
                item = _block(kind="slide_text", text=text, heading_path=heading, locator={"kind": "pptx_slide", "slide": slide_number, "shape": shape_index})
                if item: blocks.append(item)
        return blocks, {"method": "python_pptx", "slide_count": len(presentation.slides)}
    except Exception as exc:
        return [], {"status": "stored_only", "reason": f"pptx_extract_failed: {exc}"}


def _extract_xlsx(path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    try:
        from openpyxl import load_workbook  # type: ignore
        workbook = load_workbook(str(path), read_only=True, data_only=True)
        blocks = []
        sheet_count = 0
        for sheet in workbook.worksheets:
            sheet_count += 1
            rows: list[str] = []
            start_row = 1
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = ["" if value is None else str(value) for value in row]
                if not any(value.strip() for value in values):
                    continue
                if not rows:
                    start_row = row_number
                rows.append("\t".join(values))
                if len(rows) >= 40:
                    item = _block(kind="sheet_rows", text="\n".join(rows), heading_path=[sheet.title], locator={"kind": "xlsx_rows", "sheet": sheet.title, "start_row": start_row, "end_row": row_number})
                    if item: blocks.append(item)
                    rows = []
            if rows:
                item = _block(kind="sheet_rows", text="\n".join(rows), heading_path=[sheet.title], locator={"kind": "xlsx_rows", "sheet": sheet.title, "start_row": start_row, "end_row": start_row + len(rows) - 1})
                if item: blocks.append(item)
        workbook.close()
        return blocks, {"method": "openpyxl_read_only", "sheet_count": sheet_count}
    except Exception as exc:
        return [], {"status": "stored_only", "reason": f"xlsx_extract_failed: {exc}"}


def _parser_version(method: str) -> str:
    package_by_method = {
        "pypdf_all_pages": "pypdf",
        "python_docx": "python-docx",
        "python_pptx": "python-pptx",
        "openpyxl_read_only": "openpyxl",
    }
    package = package_by_method.get(str(method or ""))
    if not package:
        return "builtin"
    try:
        from importlib.metadata import version
        return str(version(package))
    except Exception:
        return "unknown"


def _source_capabilities(blocks: list[dict[str, Any]]) -> dict[str, Any]:
    locators = [dict(item.get("locator") or {}) for item in blocks if isinstance(item, dict)]
    return {
        "hierarchy": any(bool(item.get("heading_path")) for item in blocks if isinstance(item, dict)),
        "tables": any(str(item.get("kind") or "") in {"table", "table_row", "table_rows", "table_cell", "sheet_rows"} for item in blocks if isinstance(item, dict)),
        "pictures": any(str(item.get("kind") or "") in {"picture", "chart"} for item in blocks if isinstance(item, dict)),
        "page_provenance": any(locator.get("page") is not None or locator.get("kind") == "pdf_page" for locator in locators),
        "bbox_provenance": any(bool(locator.get("bbox")) for locator in locators),
        "ocr": False,
        "table_structure": any(str(item.get("kind") or "") in {"table", "table_row", "table_rows", "table_cell", "sheet_rows"} for item in blocks if isinstance(item, dict)),
    }


def _extract_project_document_lightweight(path: Path, suffix: str, *, content_hash: str = "") -> dict[str, Any]:
    if suffix == ".pdf":
        blocks, details = _extract_pdf(path)
    elif suffix == ".docx":
        blocks, details = _extract_docx(path)
    elif suffix == ".pptx":
        blocks, details = _extract_pptx(path)
    elif suffix == ".xlsx":
        blocks, details = _extract_xlsx(path)
    else:
        try:
            blocks, details = _extract_text(path, suffix)
        except Exception as exc:
            blocks, details = [], {"status": "stored_only", "reason": f"text_extract_failed: {exc}"}
    blocks = [item for item in blocks if isinstance(item, dict) and _clean(item.get("text"))]
    status = "extracted" if blocks else str(details.get("status") or "stored_only")
    method = str(details.get("method") or "neo_lightweight")
    return build_source_document(
        path=path,
        blocks=blocks,
        parser_id=method,
        parser_version=_parser_version(method),
        parser_mode="lightweight",
        status=status,
        content_hash=content_hash,
        attempts=[{"parser_id": method, "version": _parser_version(method), "status": status}],
        capabilities=_source_capabilities(blocks),
        metadata={k: v for k, v in details.items() if k not in {"status", "method"}},
        truncated=False,
    )


def extract_project_document(path: Path, suffix: str | None = None) -> dict[str, Any]:
    """Prepare a persistent project file through the ACI-2 NeoSourceDocument contract.

    Structured binary documents prefer Docling in ``auto``/``advanced`` mode.
    Neo's existing parsers remain the lightweight/fallback adapters. Set
    ``NEO_SOURCE_PREPARATION_MODE=lightweight`` to disable the advanced attempt.
    """
    path = Path(path)
    suffix = str(suffix or path.suffix).lower()
    try:
        content_hash = file_sha256(path)
    except Exception:
        content_hash = ""

    mode = str(os.environ.get("NEO_SOURCE_PREPARATION_MODE") or "auto").strip().lower()
    if mode not in {"auto", "advanced", "lightweight"}:
        mode = "auto"
    attempts: list[dict[str, Any]] = []

    if suffix in ADVANCED_DOCUMENT_SUFFIXES and mode != "lightweight":
        try:
            advanced = prepare_docling_source(path, content_hash=content_hash)
            parser = advanced.get("parser") if isinstance(advanced.get("parser"), dict) else {}
            attempts.extend(list(parser.get("attempts") or []))
            if advanced.get("blocks") and str(advanced.get("status") or "") in {"extracted", "partial"}:
                parser["attempts"] = attempts
                parser["fallback_used"] = False
                advanced["parser"] = parser
                advanced.setdefault("metadata", {})["router_mode"] = mode
                advanced["metadata"]["router_primary"] = "docling"
                return advanced
            attempts.append({"parser_id": "docling", "status": str(advanced.get("status") or "stored_only")})
        except (DoclingSourceUnavailable, DoclingSourceError) as exc:
            attempts.append({"parser_id": "docling", "status": "unavailable" if isinstance(exc, DoclingSourceUnavailable) else "error", "error": str(exc)})
        except Exception as exc:  # defensive fallback boundary
            attempts.append({"parser_id": "docling", "status": "error", "error": str(exc)})

    lightweight = _extract_project_document_lightweight(path, suffix, content_hash=content_hash)
    parser = lightweight.get("parser") if isinstance(lightweight.get("parser"), dict) else {}
    lightweight_attempts = list(parser.get("attempts") or [])
    parser["attempts"] = attempts + lightweight_attempts
    parser["fallback_used"] = bool(attempts)
    lightweight["parser"] = parser
    lightweight.setdefault("metadata", {})["router_mode"] = mode
    lightweight["metadata"]["router_primary"] = "docling" if suffix in ADVANCED_DOCUMENT_SUFFIXES and mode != "lightweight" else "lightweight"
    if attempts:
        lightweight["metadata"]["fallback_reason"] = attempts[-1].get("error") or attempts[-1].get("status") or "advanced_parser_did_not_extract"
    return lightweight


def structured_extraction_from_text(text: str, *, source_kind: str = "legacy_text") -> dict[str, Any]:
    body = str(text or "")
    blocks = _blocks_from_lines(body, markdown=source_kind in {"md", "markdown"}, source_kind=source_kind)
    return build_source_document(
        path=None,
        source_name=f"inline.{source_kind}" if source_kind else "inline_source",
        suffix="",
        blocks=blocks,
        parser_id="legacy_text_bridge",
        parser_version="builtin",
        parser_mode="lightweight",
        status="extracted" if blocks else "empty",
        content_bytes=body.encode("utf-8", errors="ignore"),
        attempts=[{"parser_id": "legacy_text_bridge", "version": "builtin", "status": "extracted" if blocks else "empty"}],
        capabilities=_source_capabilities(blocks),
        metadata={"source_kind": source_kind, "router_mode": "lightweight"},
        truncated=False,
    )


def _split_long_text(text: str, max_chars: int = MAX_FRAGMENT_CHARS, overlap: int = LONG_BLOCK_OVERLAP_CHARS) -> list[str]:
    text = _clean(text)
    if len(text) <= max_chars:
        return [text] if text else []
    pieces: list[str] = []
    start = 0
    while start < len(text):
        end = min(len(text), start + max_chars)
        if end < len(text):
            window = text[start:end]
            cut = max(window.rfind("\n\n"), window.rfind(". "), window.rfind("; "), window.rfind(" "))
            if cut >= int(max_chars * 0.6):
                end = start + cut + (1 if window[cut:cut + 1] == " " else 0)
        piece = text[start:end].strip()
        if piece:
            pieces.append(piece)
        if end >= len(text):
            break
        start = max(start + 1, end - max(0, overlap))
    return pieces


def _format_media_seconds(value: Any) -> str:
    try:
        seconds = max(0, int(float(value or 0)))
    except Exception:
        seconds = 0
    hours, rem = divmod(seconds, 3600)
    minutes, secs = divmod(rem, 60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}" if hours else f"{minutes:02d}:{secs:02d}"


def _locator_label(locator: dict[str, Any]) -> str:
    if locator.get("kind") == "media_time_range":
        if locator.get("label"):
            return str(locator.get("label"))
        return f"{_format_media_seconds(locator.get('start_seconds'))}-{_format_media_seconds(locator.get('end_seconds'))}"
    if locator.get("kind") == "pdf_page": return f"page {locator.get('page')}"
    if locator.get("kind") == "docx_paragraph": return f"paragraph {locator.get('paragraph')}"
    if locator.get("kind") == "line_range": return f"lines {locator.get('start_line')}-{locator.get('end_line')}"
    if locator.get("kind") == "pptx_slide": return f"slide {locator.get('slide')}"
    if locator.get("kind") == "xlsx_rows": return f"sheet {locator.get('sheet')} rows {locator.get('start_row')}-{locator.get('end_row')}"
    if locator.get("kind") == "source_item" and locator.get("page") is not None: return f"page {locator.get('page')}"
    if locator.get("kind") == "source_item" and locator.get("group_name"): return str(locator.get("group_name"))
    return str(locator.get("kind") or "source")


def build_contextual_fragments(extraction: dict[str, Any], *, filename: str, target_chars: int = TARGET_FRAGMENT_CHARS, max_chars: int = MAX_FRAGMENT_CHARS) -> list[dict[str, Any]]:
    raw_blocks = [item for item in (extraction.get("blocks") or []) if isinstance(item, dict) and _clean(item.get("text"))]
    parser_info = extraction.get("parser") if isinstance(extraction.get("parser"), dict) else {}
    revision_info = extraction.get("revision") if isinstance(extraction.get("revision"), dict) else {}
    source_parser = {
        "parser_id": str(parser_info.get("parser_id") or extraction.get("method") or "unknown"),
        "version": str(parser_info.get("version") or ""),
        "mode": str(parser_info.get("mode") or ""),
    }
    source_revision_id = str(revision_info.get("revision_id") or extraction.get("revision_id") or "")
    source_content_hash = str(revision_info.get("content_sha256") or extraction.get("content_hash") or "")
    expanded: list[dict[str, Any]] = []
    for block in raw_blocks:
        parts = _split_long_text(str(block.get("text") or ""), max_chars=max_chars - 500)
        for part_index, part in enumerate(parts):
            expanded.append({**block, "text": part, "split_index": part_index, "split_count": len(parts)})

    fragments: list[dict[str, Any]] = []
    current: list[dict[str, Any]] = []
    current_chars = 0
    current_heading: tuple[str, ...] = ()

    def flush() -> None:
        nonlocal current, current_chars, current_heading
        if not current:
            return
        heading_path = [item for item in current_heading if item]
        raw_text = "\n\n".join(str(item.get("text") or "") for item in current).strip()
        locators = [dict(item.get("locator") or {}) for item in current]
        location = _locator_label(locators[0]) if locators else "source"
        section = " > ".join(heading_path) if heading_path else "Document body"
        metadata = extraction.get("metadata") if isinstance(extraction.get("metadata"), dict) else {}
        source_label = str(metadata.get("source_label") or "Document").strip() or "Document"
        prefix = f"{source_label}: {filename}\nSection: {section}\nLocation: {location}\n\n"
        search_text = prefix + raw_text
        block_ids = [str(item.get("block_id") or "") for item in current]
        fragment_key = f"fragment:{_hash({'blocks': block_ids, 'text': raw_text}, 20)}"
        title = f"{filename} · {section}"
        epistemics = classify_fragment_epistemics(title=title, heading_path=heading_path, text=raw_text)
        fragments.append({
            "schema_id": PROJECT_DOCUMENT_FRAGMENT_SCHEMA_ID,
            "fragment_key": fragment_key,
            "title": title,
            "text": raw_text,
            "search_text": search_text,
            "heading_path": heading_path,
            **epistemics,
            "source_block_ids": block_ids,
            "source_locators": locators,
            "source_locator": {"kind": "structured_blocks", "blocks": locators, "label": location},
            "source_parser": dict(source_parser),
            "source_revision_id": source_revision_id,
            "source_content_hash": source_content_hash,
            "fragment_text_hash": _hash(raw_text, 64),
            "char_count": len(raw_text),
            "search_char_count": len(search_text),
        })
        current = []; current_chars = 0; current_heading = ()

    for block in expanded:
        heading = tuple(str(item) for item in (block.get("heading_path") or []) if str(item or "").strip())
        block_text = str(block.get("text") or "")
        proposed = current_chars + len(block_text) + (2 if current else 0)
        if current and (heading != current_heading or proposed > target_chars):
            flush()
        current_heading = heading
        current.append(block)
        current_chars += len(block_text) + (2 if len(current) > 1 else 0)
        if current_chars >= target_chars:
            flush()
    flush()
    return fragments


def _entity_type_from_heading(path: list[str]) -> tuple[str, str] | None:
    if len(path) < 2:
        return None
    parent = _slug(path[-2]).replace("_", " ")
    for label, entity_type in _ENTITY_SECTION_TYPES:
        if parent == label:
            return entity_type, path[-1]
    return None


def _split_aliases(value: str) -> list[str]:
    parts = re.split(r"[,;/|]", value)
    return [part.strip() for part in parts if part.strip()][:24]


def extract_explicit_knowledge(extraction: dict[str, Any]) -> dict[str, Any]:
    entities: dict[str, dict[str, Any]] = {}
    facts: list[dict[str, Any]] = []
    relationships: list[dict[str, Any]] = []

    for block in extraction.get("blocks") or []:
        if not isinstance(block, dict):
            continue
        heading_path = [str(item) for item in (block.get("heading_path") or []) if str(item or "").strip()]
        current_entity: dict[str, Any] | None = None
        inferred = _entity_type_from_heading(heading_path)
        if inferred:
            entity_type, label = inferred
            key = f"{entity_type}:{_slug(label)}"
            current_entity = entities.setdefault(key, {"entity_key": key, "entity_type": entity_type, "label": label, "aliases": [], "source_block_ids": [], "heading_path": heading_path})
            current_entity["source_block_ids"].append(block.get("block_id"))

        lines = [line.strip() for line in str(block.get("text") or "").splitlines() if line.strip()]
        for line in lines:
            explicit = _EXPLICIT_ENTITY_RE.match(line)
            if explicit:
                entity_type = _slug(explicit.group(1))
                label = explicit.group(2).strip()
                key = f"{entity_type}:{_slug(label)}"
                current_entity = entities.setdefault(key, {"entity_key": key, "entity_type": entity_type, "label": label, "aliases": [], "source_block_ids": [], "heading_path": heading_path})
                current_entity["source_block_ids"].append(block.get("block_id"))
                continue
            kv = _KEY_VALUE_RE.match(line)
            if not kv:
                continue
            key_text, value = kv.group(1).strip(), kv.group(2).strip()
            key_norm = key_text.lower().strip()
            if key_norm == "name" and value:
                entity_type = current_entity.get("entity_type") if current_entity else "entity"
                entity_key = f"{entity_type}:{_slug(value)}"
                current_entity = entities.setdefault(entity_key, {"entity_key": entity_key, "entity_type": entity_type, "label": value, "aliases": [], "source_block_ids": [], "heading_path": heading_path})
                current_entity["source_block_ids"].append(block.get("block_id"))
                continue
            if current_entity is None:
                continue
            if key_norm in _ALIAS_KEYS:
                aliases = _split_aliases(value)
                current_entity["aliases"] = sorted(set(list(current_entity.get("aliases") or []) + aliases))
                for alias in aliases:
                    facts.append({"subject_key": current_entity["entity_key"], "predicate": "identity.alias", "object_value": alias, "statement": f"{current_entity['label']} — {key_text}: {alias}", "fact_type": "source_explicit_alias", "source_block_id": block.get("block_id")})
                continue
            predicate = f"document.{_slug(key_text)}"
            facts.append({"subject_key": current_entity["entity_key"], "predicate": predicate, "object_value": value, "statement": f"{current_entity['label']} — {key_text}: {value}", "fact_type": "source_explicit_fact", "source_block_id": block.get("block_id")})
            if key_norm in _RELATION_KEYS and value and len(value) <= 180:
                relationships.append({"source_key": current_entity["entity_key"], "relation_type": f"document.{_slug(key_text)}", "target_label": value, "source_block_id": block.get("block_id")})

    return {"entities": list(entities.values()), "facts": facts, "relationships": relationships}


def enrich_fragments_with_explicit_knowledge(fragments: list[dict[str, Any]], explicit: dict[str, Any]) -> list[dict[str, Any]]:
    entities = [item for item in (explicit.get("entities") or []) if isinstance(item, dict)]
    relationships = [item for item in (explicit.get("relationships") or []) if isinstance(item, dict)]
    out: list[dict[str, Any]] = []
    for fragment in fragments:
        item = dict(fragment)
        heading = tuple(str(value) for value in (item.get("heading_path") or []) if str(value or "").strip())
        matched = None
        for entity in entities:
            entity_heading = tuple(str(value) for value in (entity.get("heading_path") or []) if str(value or "").strip())
            if entity_heading and heading[: len(entity_heading)] == entity_heading:
                matched = entity
                break
        entity_hints: list[str] = []
        relation_hints: list[str] = []
        if matched:
            label = str(matched.get("label") or "").strip()
            aliases = [str(value).strip() for value in (matched.get("aliases") or []) if str(value or "").strip()]
            if label:
                entity_hints.append(label)
            entity_hints.extend(aliases)
            source_key = str(matched.get("entity_key") or "")
            relation_hints = sorted({str(rel.get("relation_type") or "") for rel in relationships if str(rel.get("source_key") or "") == source_key and str(rel.get("relation_type") or "")})
            prefix_parts = []
            if label:
                prefix_parts.append(f"Entity: {label}")
            if aliases:
                prefix_parts.append("Aliases: " + ", ".join(aliases))
            if relation_hints:
                prefix_parts.append("Explicit relations: " + ", ".join(relation_hints))
            if prefix_parts:
                item["search_text"] = "\n".join(prefix_parts) + "\n" + str(item.get("search_text") or item.get("text") or "")
                item["search_char_count"] = len(item["search_text"])
        item["entity_hints"] = entity_hints
        item["relation_hints"] = relation_hints
        out.append(item)
    return out


def build_project_document_manifest(*, document_id: str, revision_id: str, filename: str, stored_path: str, content_hash: str, extraction: dict[str, Any], fragments: list[dict[str, Any]], explicit: dict[str, Any]) -> dict[str, Any]:
    return {
        "schema_id": PROJECT_DOCUMENT_MANIFEST_SCHEMA_ID,
        "document_id": document_id,
        "revision_id": revision_id,
        "filename": filename,
        "stored_path": stored_path,
        "content_hash": content_hash,
        "extraction": {k: v for k, v in extraction.items() if k != "blocks"},
        "fragment_count": len(fragments),
        "fragments": [{k: item.get(k) for k in ("fragment_key", "fragment_text_hash", "title", "heading_path", "source_block_ids", "source_locator", "source_parser", "source_revision_id", "source_content_hash", "char_count", "search_char_count", "epistemic_role", "epistemic_state", "claim_type", "supports_positive_claims")} for item in fragments],
        "entities": explicit.get("entities") or [],
        "facts": explicit.get("facts") or [],
        "relationships": explicit.get("relationships") or [],
        "policy": {
            "source_authority": "original_uploaded_file",
            "fragment_derivation": "direct_contextual_projection",
            "consolidation": "source_direct_no_summary_replacement",
            "canon": "source evidence is not automatically user-approved canon",
        },
    }


def projection_bundle_for_fragment(*, document_id: str, revision_id: str, filename: str, stored_path: str, fragment: dict[str, Any], context: dict[str, Any]) -> dict[str, Any]:
    fragment_key = str(fragment.get("fragment_key") or _hash(fragment, 20))
    native_id = f"{document_id}:{fragment_key}"
    native_ref = make_native_ref(
        adapter_id=PROJECT_DOCUMENT_ADAPTER_ID,
        authority_namespace="project.document.fragment",
        native_id=native_id,
        resolver_kind="project_document_fragment",
        resolver_key={
            "path": stored_path,
            "document_id": document_id,
            "revision_id": revision_id,
            "fragment_key": fragment_key,
            "fragment_text_hash": str(fragment.get("fragment_text_hash") or _hash(str(fragment.get("text") or ""), 64)),
            "source_content_hash": str(fragment.get("source_content_hash") or ""),
            "source_parser": dict(fragment.get("source_parser") or {}),
            "locator": fragment.get("source_locator") or {},
        },
        native_schema_id=PROJECT_DOCUMENT_FRAGMENT_SCHEMA_ID,
        native_revision_id=revision_id,
    )
    locator = SourceLocator(kind="project_document_fragment", path=stored_path, label=filename, fields={"document_id": document_id, "revision_id": revision_id, "fragment_key": fragment_key, "locator": fragment.get("source_locator") or {}})
    return build_projection_bundle(
        adapter_id=PROJECT_DOCUMENT_ADAPTER_ID,
        authority_namespace="project.document.fragment",
        native_id=native_id,
        kind="fragment",
        title=str(fragment.get("title") or filename),
        search_text=str(fragment.get("search_text") or fragment.get("text") or ""),
        native_ref=native_ref,
        revision_token=revision_id,
        source_locator=locator,
        context=context,
        evidence_role="project_source",
        derivation="direct_contextual_projection",
        origin="user_project_file",
        source_integrity="verified",
        claim_type=str(fragment.get("claim_type") or "source_passage"),
        epistemic_state=str(fragment.get("epistemic_state") or "established"),
        user_confirmation="provided_source",
        approval_state="not_applicable",
        canon_state="candidate",
        canon_domain_ref=str(context.get("scope_id") or context.get("project_id") or ""),
        lifecycle_state="active",
        importance="high",
        storage_class="source_authority_reference",
        entity_hints=[str(item) for item in (fragment.get("entity_hints") or [])],
        relation_hints=[str(item) for item in (fragment.get("relation_hints") or [])],
        payload={
            "document_id": document_id, "revision_id": revision_id, "fragment_key": fragment_key,
            "fragment_text_hash": str(fragment.get("fragment_text_hash") or _hash(str(fragment.get("text") or ""), 64)),
            "source_parser": dict(fragment.get("source_parser") or {}),
            "epistemic_role": str(fragment.get("epistemic_role") or "declarative_passage"),
            "supports_positive_claims": bool(fragment.get("supports_positive_claims", True)),
        },
    )


def _file_sha256(path: Path) -> str:
    return file_sha256(path)


def verify_project_document_ref(
    root_dir: Path,
    native_ref: dict[str, Any],
    *,
    expected_content_hash: str = "",
    hash_cache: dict[str, str] | None = None,
) -> dict[str, Any]:
    """Verify a previously ingested Project source without reparsing it.

    Retrieval hydration must never invoke Docling or rebuild a document merely to
    prove that an already-indexed fragment still belongs to the same source
    revision. The canonical ingestion projection already stores the fragment text.
    Foreground retrieval therefore performs only path containment + file hash
    validation and lets background/project ingestion rebuild stale projections.
    """
    root = Path(root_dir).resolve()
    ref = native_ref.get("native_ref") if isinstance(native_ref.get("native_ref"), dict) else native_ref
    resolver = ref.get("resolver_key") if isinstance(ref.get("resolver_key"), dict) else {}
    path_text = str(resolver.get("path") or "").strip()
    if not path_text:
        return {"ok": False, "status": "missing", "source_integrity": "missing", "adapter_id": PROJECT_DOCUMENT_ADAPTER_ID, "reason": "missing_path"}
    path = Path(path_text)
    if not path.is_absolute():
        path = root / path
    try:
        path = path.resolve()
    except Exception:
        return {"ok": False, "status": "missing", "source_integrity": "missing", "adapter_id": PROJECT_DOCUMENT_ADAPTER_ID, "reason": "invalid_path"}
    if root != path and root not in path.parents:
        return {"ok": False, "status": "unavailable", "source_integrity": "unavailable", "adapter_id": PROJECT_DOCUMENT_ADAPTER_ID, "reason": "path_outside_root"}
    if not path.exists() or not path.is_file():
        return {"ok": False, "status": "missing", "source_integrity": "missing", "adapter_id": PROJECT_DOCUMENT_ADAPTER_ID, "path": str(path)}

    cache = hash_cache if isinstance(hash_cache, dict) else {}
    try:
        stat = path.stat()
        cache_key = f"{path}|{stat.st_size}|{getattr(stat, 'st_mtime_ns', int(stat.st_mtime * 1e9))}"
    except Exception:
        cache_key = str(path)
    current_hash = cache.get(cache_key)
    if not current_hash:
        current_hash = _file_sha256(path)
        cache[cache_key] = current_hash

    expected_revision = str(resolver.get("revision_id") or ref.get("native_revision_id") or "")
    expected_hash = str(expected_content_hash or resolver.get("source_content_hash") or "").strip()
    current_revision = f"rev_{current_hash[:24]}"
    revision_matches = not expected_revision or expected_revision == current_revision
    hash_matches = not expected_hash or expected_hash == current_hash
    matches = bool(revision_matches and hash_matches)
    return {
        "ok": matches,
        "status": "verified" if matches else "stale",
        "source_integrity": "verified" if matches else "stale",
        "adapter_id": PROJECT_DOCUMENT_ADAPTER_ID,
        "path": str(path),
        "expected_revision_id": expected_revision,
        "current_revision_id": current_revision,
        "expected_content_hash": expected_hash,
        "current_content_hash": current_hash,
        "verification_mode": "stored_projection_hash_only",
        "reparse_performed": False,
        "reason": "" if matches else "source_revision_changed_reingest_required",
    }


def resolve_project_document_ref(root_dir: Path, native_ref: dict[str, Any]) -> dict[str, Any]:
    """Hydrate an NKB-7 project fragment from the current original source file.

    This is deliberately independent from the NKB-6 native adapter registry. Public
    project files are user-supplied authorities, not Neo-native stores. NKB-8 can
    call this resolver when authority adjudication needs the current source passage.
    """
    root = Path(root_dir).resolve()
    ref = native_ref.get("native_ref") if isinstance(native_ref.get("native_ref"), dict) else native_ref
    resolver = ref.get("resolver_key") if isinstance(ref.get("resolver_key"), dict) else {}
    path_text = str(resolver.get("path") or "").strip()
    if not path_text:
        return {"ok": False, "status": "missing", "adapter_id": PROJECT_DOCUMENT_ADAPTER_ID, "reason": "missing_path"}
    path = Path(path_text)
    if not path.is_absolute():
        path = root / path
    try:
        path = path.resolve()
    except Exception:
        return {"ok": False, "status": "missing", "adapter_id": PROJECT_DOCUMENT_ADAPTER_ID, "reason": "invalid_path"}
    if root != path and root not in path.parents:
        return {"ok": False, "status": "unavailable", "adapter_id": PROJECT_DOCUMENT_ADAPTER_ID, "reason": "path_outside_root"}
    if not path.exists() or not path.is_file():
        return {"ok": False, "status": "missing", "adapter_id": PROJECT_DOCUMENT_ADAPTER_ID, "path": str(path)}

    current_hash = _file_sha256(path)
    expected_revision = str(resolver.get("revision_id") or ref.get("native_revision_id") or "")
    current_revision = f"rev_{current_hash[:24]}"
    extraction = extract_project_document(path, path.suffix.lower())
    fragments = build_contextual_fragments(extraction, filename=path.name)
    explicit = extract_explicit_knowledge(extraction)
    fragments = enrich_fragments_with_explicit_knowledge(fragments, explicit)
    wanted = str(resolver.get("fragment_key") or "")
    wanted_text_hash = str(resolver.get("fragment_text_hash") or "")
    fragment = next((item for item in fragments if str(item.get("fragment_key") or "") == wanted), None)
    if fragment is None and wanted_text_hash:
        fragment = next((item for item in fragments if str(item.get("fragment_text_hash") or _hash(str(item.get("text") or ""), 64)) == wanted_text_hash), None)
    revision_matches = not expected_revision or expected_revision == current_revision
    current_parser = extraction.get("parser") if isinstance(extraction.get("parser"), dict) else {}
    expected_parser = resolver.get("source_parser") if isinstance(resolver.get("source_parser"), dict) else {}
    parser_changed = bool(expected_parser) and (
        str(expected_parser.get("parser_id") or "") != str(current_parser.get("parser_id") or extraction.get("method") or "")
        or (str(expected_parser.get("version") or "") and str(expected_parser.get("version") or "") != str(current_parser.get("version") or ""))
    )
    if fragment is None:
        return {
            "ok": False,
            "status": "stale" if not revision_matches else ("unavailable" if parser_changed else "missing"),
            "adapter_id": PROJECT_DOCUMENT_ADAPTER_ID,
            "path": str(path),
            "expected_revision_id": expected_revision,
            "current_revision_id": current_revision,
            "expected_parser": expected_parser,
            "current_parser": current_parser,
            "reason": "parser_changed_fragment_not_resolved" if revision_matches and parser_changed else "fragment_not_found_in_current_source",
        }
    return {
        "ok": True,
        "status": "verified" if revision_matches else "stale",
        "adapter_id": PROJECT_DOCUMENT_ADAPTER_ID,
        "path": str(path),
        "document_id": str(resolver.get("document_id") or ""),
        "expected_revision_id": expected_revision,
        "current_revision_id": current_revision,
        "fragment": fragment,
        "text": fragment.get("text") or "",
        "search_text": fragment.get("search_text") or "",
        "source_locator": fragment.get("source_locator") or {},
        "source_integrity": "verified" if revision_matches else "stale",
        "expected_parser": expected_parser,
        "current_parser": current_parser,
        "parser_changed": parser_changed,
    }
