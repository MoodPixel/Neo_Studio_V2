from __future__ import annotations

import json
import math
import textwrap
from pathlib import Path
from typing import Any

from .contracts import ACI1_MANIFEST_SCHEMA_ID


def _font_path() -> str | None:
    candidates = [
        Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'),
        Path('/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf'),
    ]
    for candidate in candidates:
        if candidate.exists():
            return str(candidate)
    return None


def _draw_wrapped(draw: Any, xy: tuple[int, int], text: str, *, font: Any, fill: Any = 'black', width: int = 72, line_gap: int = 8) -> int:
    x, y = xy
    for paragraph in text.split('\n'):
        lines = textwrap.wrap(paragraph, width=max(10, width)) or ['']
        for line in lines:
            draw.text((x, y), line, font=font, fill=fill)
            bbox = draw.textbbox((x, y), line or 'Ag', font=font)
            y += max(18, int(bbox[3] - bbox[1]) + line_gap)
        y += line_gap
    return y


def _make_docx(path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_BREAK

    doc = Document()
    doc.add_heading('Crownport Operations Manual', 0)
    doc.add_paragraph('ACI1_DOCX_EARLY_FACT: Meridian intake opens at 07:30.')
    doc.add_heading('Operating Principles', level=1)
    doc.add_paragraph('This reference fixture deliberately contains hierarchy, tables, repeated prose, and a late fact beyond the ordinary Assistant extraction budget.')
    table = doc.add_table(rows=4, cols=3)
    table.style = 'Table Grid'
    cells = [
        ('Team', 'Window', 'Code'),
        ('North', '07:30', 'M-17'),
        ('South', '08:10', 'S-22'),
        ('Archive', '18:45', 'A-09'),
    ]
    for r, row in enumerate(cells):
        for c, value in enumerate(row):
            table.cell(r, c).text = value

    filler = (
        'Operational note: staff verify source identity, preserve revision provenance, and record every handoff before the next processing stage. '
        'This sentence is intentionally repeated to create a realistically large structured Word document without using private or copyrighted material.'
    )
    for section in range(1, 13):
        doc.add_heading(f'Section {section}: Controlled Procedure', level=1)
        for paragraph in range(1, 18):
            doc.add_paragraph(f'{section}.{paragraph} {filler}')
        if section % 3 == 0:
            p = doc.add_paragraph()
            p.add_run().add_break(WD_BREAK.PAGE)

    doc.add_heading('Restricted Appendix', level=1)
    doc.add_heading('Lantern Procedure', level=2)
    doc.add_paragraph('ACI1_DOCX_LATE_FACT: Project Lantern uses vault code ORCHID-731.')
    doc.add_paragraph('ACI1_DOCX_LATE_OWNER: The designated custodian is Rowan Vale.')
    doc.save(path)


def _pdf_canvas(path: Path, pagesize: tuple[float, float] | None = None):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    return canvas.Canvas(str(path), pagesize=pagesize or A4)


def _make_long_pdf(path: Path) -> None:
    from reportlab.lib.pagesizes import A4

    c = _pdf_canvas(path, A4)
    width, height = A4
    for page in range(1, 121):
        c.setFont('Helvetica-Bold', 16)
        c.drawString(54, height - 64, f'ACI-1 Long PDF — Page {page}')
        c.setFont('Helvetica', 10)
        y = height - 96
        for row in range(15):
            c.drawString(54, y, f'Page {page} control paragraph {row + 1}: source provenance remains bound to the canonical record.')
            y -= 22
        if page == 75:
            c.setFont('Helvetica-Bold', 12)
            c.drawString(54, 180, 'ACI1_PDF_PAGE75_SECRET: HELIOTROPE-75')
            c.drawString(54, 158, 'The recovery locker is located in Bay C-14.')
        c.setFont('Helvetica', 8)
        c.drawRightString(width - 54, 32, f'{page}')
        c.showPage()
    c.save()


def _make_multicolumn_pdf(path: Path) -> None:
    from reportlab.lib.pagesizes import A4

    c = _pdf_canvas(path, A4)
    width, height = A4
    margin = 48
    gutter = 24
    col_width = (width - (2 * margin) - gutter) / 2
    body = [
        'Context preparation should preserve reading order and source locality rather than treating every extracted token as a flat bag of text.',
        'The benchmark repeats realistic prose so each column has enough material to expose ordering mistakes while remaining fully redistributable.',
    ]
    for page in range(1, 4):
        c.setFont('Helvetica-Bold', 15)
        c.drawString(margin, height - 52, f'Multi-column Brief — Page {page}')
        for column in range(2):
            x = margin + column * (col_width + gutter)
            y = height - 82
            c.setFont('Helvetica-Bold', 10)
            c.drawString(x, y, 'LEFT STREAM' if column == 0 else 'RIGHT STREAM')
            y -= 20
            c.setFont('Helvetica', 9)
            for idx in range(12):
                words = f'{page}.{column + 1}.{idx + 1} {body[(idx + column) % len(body)]}'
                for line in textwrap.wrap(words, width=48):
                    c.drawString(x, y, line)
                    y -= 12
                y -= 5
            if page == 2 and column == 1:
                c.setFont('Helvetica-Bold', 10)
                c.drawString(x, max(72, y - 8), 'ACI1_MULTI_COLUMN_RIGHT: COBALT-COMPASS')
        c.showPage()
    c.save()


def _make_scanned_pdf(path: Path, asset_dir: Path) -> None:
    from PIL import Image, ImageDraw, ImageFont

    font_path = _font_path()
    font = ImageFont.truetype(font_path, 32) if font_path else ImageFont.load_default()
    small = ImageFont.truetype(font_path, 24) if font_path else ImageFont.load_default()
    images: list[Image.Image] = []
    for page in range(1, 4):
        img = Image.new('RGB', (1654, 2339), 'white')
        draw = ImageDraw.Draw(img)
        draw.text((110, 110), f'Scanned Operations Note — Page {page}', font=font, fill='black')
        y = 190
        for idx in range(18):
            y = _draw_wrapped(
                draw,
                (110, y),
                f'Scan line {idx + 1}. This page is raster-only and contains no PDF text layer.',
                font=small,
                width=72,
                line_gap=4,
            )
        if page == 2:
            draw.rectangle((95, 1580, 1550, 1780), outline='black', width=4)
            _draw_wrapped(draw, (125, 1610), 'ACI1_SCAN_PAGE2_SECRET: VIOLET-OWL-204\nEmergency cabinet: East Hall 9.', font=font, width=52)
        page_path = asset_dir / f'scan_page_{page}.png'
        img.save(page_path)
        images.append(img)
    first, rest = images[0], images[1:]
    first.save(path, 'PDF', resolution=150.0, save_all=True, append_images=rest)


def _make_table_pdf(path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    story: list[Any] = [Paragraph('Inventory Reconciliation Tables', styles['Title']), Spacer(1, 6 * mm)]
    for page in range(1, 5):
        story.append(Paragraph(f'Table Group {page}', styles['Heading1']))
        data = [['Item', 'Zone', 'Units', 'Verification']]
        for row in range(1, 26):
            verification = f'CHK-{page:02d}-{row:02d}'
            units = page * 100 + row
            if page == 3 and row == 17:
                verification = 'ACI1_TABLE_CELL_X9: 4827'
                units = 4827
            data.append([f'Part-{page}-{row:02d}', f'Z{(row % 6) + 1}', str(units), verification])
        table = Table(data, colWidths=[42 * mm, 30 * mm, 28 * mm, 70 * mm], repeatRows=1)
        table.setStyle(TableStyle([
            ('GRID', (0, 0), (-1, -1), 0.35, colors.black),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 7),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ]))
        story.append(table)
        if page != 4:
            story.append(PageBreak())
    SimpleDocTemplate(str(path), pagesize=A4, rightMargin=16 * mm, leftMargin=16 * mm, topMargin=14 * mm, bottomMargin=14 * mm).build(story)


def _make_chart_pdf(path: Path, asset_dir: Path) -> None:
    from PIL import Image, ImageDraw, ImageFont
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader

    font_path = _font_path()
    font = ImageFont.truetype(font_path, 28) if font_path else ImageFont.load_default()
    small = ImageFont.truetype(font_path, 22) if font_path else ImageFont.load_default()
    chart = Image.new('RGB', (1200, 700), 'white')
    draw = ImageDraw.Draw(chart)
    draw.line((100, 580, 1100, 580), fill='black', width=4)
    draw.line((100, 80, 100, 580), fill='black', width=4)
    values = [180, 330, 250, 470, 390]
    labels = ['A', 'B', 'C', 'D', 'E']
    for i, value in enumerate(values):
        x = 170 + i * 185
        draw.rectangle((x, 580 - value, x + 95, 580), outline='black', width=3)
        draw.text((x + 22, 600), labels[i], font=small, fill='black')
        draw.text((x + 8, 535 - value), str(value), font=small, fill='black')
    draw.text((110, 20), 'Quarterly Signal Index', font=font, fill='black')
    draw.text((630, 110), 'ACI1_CHART_IMAGE_MARKER: AMBER-DELTA-9', font=small, fill='black')
    chart_path = asset_dir / 'signal_chart.png'
    chart.save(chart_path)

    c = _pdf_canvas(path, A4)
    width, height = A4
    c.setFont('Helvetica-Bold', 17)
    c.drawString(54, height - 60, 'Visual Analysis Report')
    c.setFont('Helvetica', 10)
    c.drawString(54, height - 84, 'Figure 1 summarizes the synthetic signal index used by the ACI-1 reference corpus.')
    c.drawImage(ImageReader(chart), 54, 235, width=width - 108, height=330, preserveAspectRatio=True, anchor='c')
    c.setFont('Helvetica-Bold', 10)
    c.drawString(54, 205, 'Figure 1 caption: Signal index peak occurs in category D.')
    c.setFont('Helvetica', 10)
    c.drawString(54, 183, 'ACI1_CHART_CAPTION_FACT: category D is the highest plotted value.')
    c.showPage()
    c.save()


def _make_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    for slide_no in range(1, 7):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f'Operations Brief — Slide {slide_no}'
        body = slide.placeholders[1].text_frame
        body.clear()
        p = body.paragraphs[0]
        p.text = f'Slide {slide_no} contains structured presentation content for the ACI-1 benchmark.'
        p.font.size = Pt(22)
        for idx in range(3):
            p = body.add_paragraph()
            p.text = f'Point {idx + 1}: preserve slide hierarchy and speaker-facing reading order.'
            p.level = 0
        if slide_no == 4:
            box = slide.shapes.add_textbox(Inches(1), Inches(5.6), Inches(8), Inches(0.8))
            box.text_frame.text = 'ACI1_PPTX_SLIDE4_SECRET: NEON-FALCON-44'
    prs.save(path)


def _make_xlsx(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    summary = wb.active
    summary.title = 'Summary'
    summary.append(['Metric', 'Q1', 'Q2', 'Q3', 'Q4'])
    for row in range(1, 12):
        summary.append([f'Metric {row}', row * 11, row * 13, row * 17, row * 19])
    for cell in summary[1]:
        cell.font = Font(bold=True)

    costs = wb.create_sheet('Costs')
    costs.append(['Account', 'Owner', 'Region', 'Amount'])
    for row in range(2, 26):
        costs.cell(row=row, column=1, value=f'ACCT-{row:03d}')
        costs.cell(row=row, column=2, value=f'Owner {row - 1}')
        costs.cell(row=row, column=3, value=['North', 'South', 'East', 'West'][row % 4])
        costs.cell(row=row, column=4, value=round(100 + math.sin(row) * 25 + row * 3.2, 2))
    costs['D18'] = 913.42
    costs['E18'] = 'ACI1_XLSX_CELL_D18: 913.42'
    wb.save(path)


def _make_markdown(path: Path) -> None:
    path.write_text(
        '''# ACI-1 Markdown Baseline\n\nThis source should already work well through Neo's legacy text path.\n\n## Retrieval Contract\n\nACI1_MARKDOWN_FACT: The baseline beacon is SILVER-PINE-18.\n\n### Notes\n\n- preserve headings\n- preserve exact text\n- avoid unnecessary parser overhead\n''',
        encoding='utf-8',
    )


def _make_html(path: Path) -> None:
    path.write_text(
        '''<!doctype html><html><head><title>ACI-1 HTML Fixture</title></head><body><main><h1>Dispatch Notes</h1><section><h2>Primary Route</h2><p>The normal dispatch route uses Gate 7.</p></section><section><h2>Fallback Route</h2><table><tr><th>Route</th><th>Code</th></tr><tr><td>Night</td><td>ACI1_HTML_TABLE_FACT: INDIGO-72</td></tr></table></section></main></body></html>''',
        encoding='utf-8',
    )


def _make_screenshot(path: Path) -> None:
    from PIL import Image, ImageDraw, ImageFont

    font_path = _font_path()
    title = ImageFont.truetype(font_path, 36) if font_path else ImageFont.load_default()
    body = ImageFont.truetype(font_path, 25) if font_path else ImageFont.load_default()
    img = Image.new('RGB', (1500, 980), 'white')
    draw = ImageDraw.Draw(img)
    draw.rounded_rectangle((70, 70, 1430, 910), radius=24, outline='black', width=4)
    draw.text((115, 115), 'Incident Dashboard Screenshot', font=title, fill='black')
    y = 200
    lines = [
        'Status: Review required',
        'Queue: Context Preparation',
        'ACI1_SCREENSHOT_SECRET: CRIMSON-LANTERN-6',
        'Assigned bay: West 12',
        'Note: this fixture is a raster document image with no embedded text layer.',
    ]
    for line in lines:
        draw.text((125, y), line, font=body, fill='black')
        y += 90
    img.save(path)


def _make_messy_pdf(path: Path) -> None:
    from reportlab.lib.pagesizes import A4

    c = _pdf_canvas(path, A4)
    width, height = A4
    c.setFont('Helvetica-Bold', 16)
    c.drawString(48, height - 48, 'Mixed Layout Field Report')
    c.setFont('Helvetica', 9)
    # left narrative
    y = height - 82
    for idx in range(14):
        c.drawString(48, y, f'Left narrative {idx + 1}: inspection notes continue across the page.')
        y -= 16
    # right callout
    c.rect(width / 2 + 14, height - 290, width / 2 - 62, 180)
    c.setFont('Helvetica-Bold', 10)
    c.drawString(width / 2 + 28, height - 136, 'SIDEBAR / EXCEPTION')
    c.setFont('Helvetica', 9)
    c.drawString(width / 2 + 28, height - 160, 'ACI1_MESSY_SIDEBAR_FACT: TEAL-RAVEN-88')
    c.drawString(width / 2 + 28, height - 180, 'Route exception: service tunnel B.')
    # mini table
    start_y = 300
    c.setFont('Helvetica-Bold', 9)
    headers = ['Node', 'State', 'Owner']
    xs = [48, 180, 320]
    for x, text in zip(xs, headers):
        c.drawString(x, start_y, text)
    c.setFont('Helvetica', 9)
    for r in range(1, 7):
        row_y = start_y - r * 22
        vals = [f'N-{r}', 'Ready' if r % 2 else 'Hold', f'Team {r}']
        for x, text in zip(xs, vals):
            c.drawString(x, row_y, text)
    c.saveState()
    c.translate(width - 26, 90)
    c.rotate(90)
    c.setFont('Helvetica', 8)
    c.drawString(0, 0, 'Rotated margin note: preserve geometry when possible.')
    c.restoreState()
    c.showPage()
    c.save()


def _fixture(id_: str, path: Path, category: str, required_terms: list[str], queries: list[dict[str, Any]], *, parity: bool = True, note: str = '') -> dict[str, Any]:
    return {
        'id': id_,
        'path': str(path),
        'category': category,
        'parity': parity,
        'note': note,
        'truth': {
            'required_terms': required_terms,
            'queries': queries,
        },
    }


def build_reference_corpus(root: Path) -> dict[str, Any]:
    root = Path(root).resolve()
    fixtures_dir = root / 'fixtures'
    assets_dir = root / 'assets'
    fixtures_dir.mkdir(parents=True, exist_ok=True)
    assets_dir.mkdir(parents=True, exist_ok=True)

    paths = {
        'structured_docx': fixtures_dir / 'structured_large.docx',
        'long_pdf': fixtures_dir / 'long_120_pages.pdf',
        'multi_column_pdf': fixtures_dir / 'multi_column.pdf',
        'scanned_pdf': fixtures_dir / 'scanned_image_only.pdf',
        'table_pdf': fixtures_dir / 'table_heavy.pdf',
        'chart_pdf': fixtures_dir / 'image_chart_heavy.pdf',
        'pptx': fixtures_dir / 'structured_deck.pptx',
        'xlsx': fixtures_dir / 'structured_workbook.xlsx',
        'markdown': fixtures_dir / 'baseline.md',
        'html': fixtures_dir / 'structured.html',
        'screenshot': fixtures_dir / 'document_screenshot.png',
        'messy_pdf': fixtures_dir / 'messy_mixed_layout.pdf',
    }

    _make_docx(paths['structured_docx'])
    _make_long_pdf(paths['long_pdf'])
    _make_multicolumn_pdf(paths['multi_column_pdf'])
    _make_scanned_pdf(paths['scanned_pdf'], assets_dir)
    _make_table_pdf(paths['table_pdf'])
    _make_chart_pdf(paths['chart_pdf'], assets_dir)
    _make_pptx(paths['pptx'])
    _make_xlsx(paths['xlsx'])
    _make_markdown(paths['markdown'])
    _make_html(paths['html'])
    _make_screenshot(paths['screenshot'])
    _make_messy_pdf(paths['messy_pdf'])

    fixtures = [
        _fixture(
            'structured_large_docx', paths['structured_docx'], 'structured_docx',
            ['ACI1_DOCX_EARLY_FACT', 'ACI1_DOCX_LATE_FACT', 'ORCHID-731'],
            [{'id': 'docx_late_fact', 'query': 'What vault code does Project Lantern use?', 'must_retrieve_any': ['ORCHID-731', 'ACI1_DOCX_LATE_FACT'], 'answer_must_contain_any': ['ORCHID-731']}],
        ),
        _fixture(
            'long_pdf_after_page_40', paths['long_pdf'], 'long_pdf',
            ['ACI1_PDF_PAGE75_SECRET', 'HELIOTROPE-75', 'Bay C-14'],
            [{'id': 'page75_secret', 'query': 'What is the secret on page 75 and where is the recovery locker?', 'must_retrieve_all': ['HELIOTROPE-75', 'Bay C-14'], 'answer_must_contain_all': ['HELIOTROPE-75', 'Bay C-14'], 'expected_locator': {'page': 75}}],
        ),
        _fixture(
            'multi_column_pdf', paths['multi_column_pdf'], 'multi_column_pdf',
            ['ACI1_MULTI_COLUMN_RIGHT', 'COBALT-COMPASS'],
            [{'id': 'right_column', 'query': 'What marker appears in the right stream on page 2?', 'must_retrieve_any': ['COBALT-COMPASS'], 'answer_must_contain_any': ['COBALT-COMPASS'], 'expected_locator': {'page': 2}}],
        ),
        _fixture(
            'scanned_image_only_pdf', paths['scanned_pdf'], 'scanned_pdf',
            ['ACI1_SCAN_PAGE2_SECRET', 'VIOLET-OWL-204', 'East Hall 9'],
            [{'id': 'scan_secret', 'query': 'What secret appears on the scanned second page?', 'must_retrieve_all': ['VIOLET-OWL-204', 'East Hall 9'], 'answer_must_contain_all': ['VIOLET-OWL-204', 'East Hall 9'], 'expected_locator': {'page': 2}}],
        ),
        _fixture(
            'table_heavy_pdf', paths['table_pdf'], 'table_pdf',
            ['ACI1_TABLE_CELL_X9', '4827'],
            [{'id': 'table_cell', 'query': 'What verification value is in the special inventory table cell?', 'must_retrieve_any': ['ACI1_TABLE_CELL_X9', '4827'], 'answer_must_contain_any': ['4827'], 'expected_locator': {'page': 3}}],
        ),
        _fixture(
            'image_chart_heavy_pdf', paths['chart_pdf'], 'chart_pdf',
            ['ACI1_CHART_CAPTION_FACT', 'category D', 'ACI1_CHART_IMAGE_MARKER'],
            [
                {'id': 'chart_peak', 'query': 'Which category has the highest plotted signal?', 'must_retrieve_all': ['category D', 'ACI1_CHART_CAPTION_FACT'], 'answer_must_contain_any': ['category D'], 'expected_locator': {'page': 1}},
                {'id': 'chart_image_marker', 'query': 'What marker is printed inside the chart image?', 'must_retrieve_any': ['ACI1_CHART_IMAGE_MARKER', 'AMBER-DELTA-9'], 'answer_must_contain_any': ['AMBER-DELTA-9'], 'expected_locator': {'page': 1}},
            ],
        ),
        _fixture(
            'pptx_structured_deck', paths['pptx'], 'pptx',
            ['ACI1_PPTX_SLIDE4_SECRET', 'NEON-FALCON-44'],
            [{'id': 'slide4_secret', 'query': 'What secret marker is on slide 4?', 'must_retrieve_any': ['NEON-FALCON-44', 'ACI1_PPTX_SLIDE4_SECRET'], 'answer_must_contain_any': ['NEON-FALCON-44']}],
        ),
        _fixture(
            'xlsx_structured_workbook', paths['xlsx'], 'xlsx',
            ['ACI1_XLSX_CELL_D18', '913.42'],
            [{'id': 'xlsx_cell', 'query': 'What amount is recorded in the marked Costs row?', 'must_retrieve_all': ['ACI1_XLSX_CELL_D18', '913.42'], 'answer_must_contain_any': ['913.42']}],
        ),
        _fixture(
            'markdown_baseline', paths['markdown'], 'markdown',
            ['ACI1_MARKDOWN_FACT', 'SILVER-PINE-18'],
            [{'id': 'markdown_beacon', 'query': 'What is the baseline beacon?', 'must_retrieve_any': ['SILVER-PINE-18'], 'answer_must_contain_any': ['SILVER-PINE-18']}],
        ),
        _fixture(
            'html_structured', paths['html'], 'html',
            ['ACI1_HTML_TABLE_FACT', 'INDIGO-72'],
            [{'id': 'html_route', 'query': 'What code is listed for the night fallback route?', 'must_retrieve_any': ['INDIGO-72'], 'answer_must_contain_any': ['INDIGO-72']}],
        ),
        _fixture(
            'document_screenshot', paths['screenshot'], 'document_screenshot',
            ['ACI1_SCREENSHOT_SECRET', 'CRIMSON-LANTERN-6'],
            [{'id': 'screenshot_secret', 'query': 'What secret appears in the dashboard screenshot?', 'must_retrieve_any': ['CRIMSON-LANTERN-6'], 'answer_must_contain_any': ['CRIMSON-LANTERN-6']}],
            parity=False,
            note='Supplemental document-image preparation fixture. Current Neo image attachments use native image intelligence rather than the legacy document extractor, so this row must not be used alone to claim a production regression.',
        ),
        _fixture(
            'messy_mixed_layout_pdf', paths['messy_pdf'], 'messy_mixed_pdf',
            ['ACI1_MESSY_SIDEBAR_FACT', 'TEAL-RAVEN-88', 'service tunnel B'],
            [{'id': 'messy_sidebar', 'query': 'What route exception is shown in the sidebar?', 'must_retrieve_all': ['TEAL-RAVEN-88', 'service tunnel B'], 'answer_must_contain_all': ['TEAL-RAVEN-88', 'service tunnel B'], 'expected_locator': {'page': 1}}],
        ),
    ]

    manifest_path = root / 'fixture_manifest.generated.json'
    manifest = {
        'schema_id': ACI1_MANIFEST_SCHEMA_ID,
        'description': 'Deterministic, redistributable ACI-1B reference corpus generated by Neo. Truth markers are authored from the generator source, not from parser output.',
        'corpus_kind': 'neo_generated_reference',
        'private_data': False,
        'fixtures': fixtures,
    }
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding='utf-8')

    provenance = {
        'schema_id': 'neo.aci1.reference_corpus.provenance.v1',
        'generator': 'neo_app.experiments.aci1.reference_corpus.build_reference_corpus',
        'redistribution': 'Neo-generated synthetic benchmark content; no private/client material and no third-party copyrighted source text.',
        'fixture_count': len(fixtures),
        'manifest': str(manifest_path),
    }
    (root / 'CORPUS_PROVENANCE.json').write_text(json.dumps(provenance, indent=2), encoding='utf-8')
    return {'root': str(root), 'manifest': str(manifest_path), 'fixture_count': len(fixtures), 'fixtures': fixtures}
